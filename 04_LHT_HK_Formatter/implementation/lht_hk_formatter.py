"""Create a standardized six-sheet LHT cost workbook from SAP Excel exports.

Requires Python 3.10+, openpyxl, pandas.
Usage:
    pip install openpyxl pandas
    python lht_hk_formatter.py source.xlsx -o output.xlsx

SAP exports differ by transaction and configuration. This implementation uses
conservative heuristics and reports uncertainty instead of inventing data.
"""
from __future__ import annotations
import argparse, math, os, re, json
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

NAVY, WHITE, GOLD, AMBER = "0D1B2A", "FFFFFF", "FFD966", "F4B183"
LIGHT_BLUE, LIGHT_GREY, YELLOW = "DDEBF7", "E7E6E6", "FFF2CC"
SHEETS = ["Cover", "Assembly", "BOM_Detail", "Pivot_Assembly", "Operations", "Suppliers"]
ASSEMBLY = ["Pos","HkGr","Material","Bezeichnung","Menge","Einheit","Einzelwert_EUR","Gesamtwert_EUR","Einzel_Sammel","Waehrung","Row_Type","BOM_Level","Level_Label","Parent_Path"]
BOM = ["BOM_Level","Level_Label","Pos","HkGr","Material","Bezeichnung","Ressource","Werk","Material_No","Operation","Menge","Einheit","Einzelwert_EUR","Gesamtwert_EUR","Einzel_Sammel","Waehrung","Row_Type","Op_Variant","Parent_Path"]
PIVOT = ["HkGr","Count","Total_EUR","Share_%"]
OPS = ["Pos","HkGr","ArbPlatz","Bezeichnung","Menge","Einheit","Einzelwert_EUR","Gesamtwert_EUR","STD_Hours","BOM_Level","Parent_Path"]
SUPPLIERS = ["Material","Bezeichnung","Kreditor","Vendor_Name","Menge","Einheit","Einzelwert_EUR","Gesamtwert_EUR","HkGr"]
ALIASES = {
"Pos":["pos","position","item","item no","item number"],
"HkGr":["hkgr","hk-gr","kostenartengruppe","cost group"],
"Material":["material","materialnummer","material number","matnr","material no"],
"Bezeichnung":["bezeichnung","materialbeschreibung","beschreibung","text","kurztext","description","operation text"],
"Menge":["menge","quantity","qty","basismenge","bedarf"],
"Einheit":["einheit","me","unit","mengeneinheit"],
"Einzelwert_EUR":["einzelwert","einzelkosten","preis je einheit","unit cost","single value"],
"Gesamtwert_EUR":["gesamtwert","gesamtkosten","wert gesamt","total value","total cost","kosten"],
"Einzel_Sammel":["einzel/sammel","einzel_sammel","einzel-sammel"],
"Waehrung":["waehrung","währung","currency","curr"],
"Ressource":["ressource","resource","ressourcen"],
"Werk":["werk","plant"],
"Material_No":["material_no","material no","materialnummer","matnr"],
"Operation":["operation","vorgang","arbeitsgang"],
"ArbPlatz":["arbplatz","arbeitsplatz","work center"],
"STD_Hours":["std_hours","std hours","stunden","arbeitszeit","activity quantity","std menge"],
"Kostenstelle":["kostenstelle","cost center","kst"],
"Leistungsart":["leistungsart","activity type"],
"Kreditor":["kreditor","lieferant","vendor","vendor number","lieferantennummer","lifnr"],
"Vendor_Name":["vendor_name","lieferantenname","vendor name"],
"BOM_Level":["bom_level","ebene","stufe","hierarchiestufe"],
"Op_Variant":["op_variant","vorgangsart","operation variant"]
}
def norm(v: Any) -> str:
    return re.sub(r"\s+"," ",str(v or "").strip().lower())
def parse_number(v: Any):
    if v is None or (isinstance(v,float) and math.isnan(v)): return None
    if isinstance(v,(int,float)) and not isinstance(v,bool): return float(v)
    s=str(v).strip()
    if not s or s in {"-","—","–","n/a","N/A"}: return None
    neg=s.startswith("(") and s.endswith(")")
    if neg: s=s[1:-1]
    s=s.replace("€","").replace("EUR","").replace("eur","").replace("\u00a0","").replace(" ","")
    if "," in s: s=s.replace(".","").replace(",",".")
    elif re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+",s): s=s.replace(".","")
    try:
        x=float(s); return -x if neg else x
    except (ValueError,TypeError): return v

def find_header(ws):
    best=(0,0,1,[])
    aliases=[norm(a) for vals in ALIASES.values() for a in vals]
    for r in range(1,min(30,ws.max_row)+1):
        vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
        score=sum(1 for v in vals if v is not None and any(norm(v)==a or (len(a)>3 and a in norm(v)) for a in aliases))
        nonempty=sum(1 for v in vals if v is not None and str(v).strip())
        if (score,nonempty)>(best[0],best[1]): best=(score,nonempty,r,vals)
    return best[2],best[3],best[0]

def map_columns(headers):
    result={}
    for i,h in enumerate(headers):
        nh=norm(h)
        for key,aliases in ALIASES.items():
            if nh in [norm(a) for a in aliases] or any(len(norm(a))>3 and norm(a) in nh for a in aliases):
                result.setdefault(key,i)
    return result

def detect_tx(text, headers):
    t=norm(text)
    for tx in ["CK13N","CK53N","ZCO6","ZCO9"]:
        if norm(tx) in t:
            mode={"CK13N":"full_bom","CK53N":"flat","ZCO6":"period_comparison","ZCO9":"auto"}[tx]
            return tx,mode,f"{tx} found in header/first 19 rows"
    cols=" ".join(map(norm,headers))
    if "ressource" in cols or "hierarchiestufe" in cols or "parent" in cols:
        return "Unknown","full_bom","Inferred from hierarchy/resource-like columns"
    return "Unknown","flat","Inferred from itemized column structure; transaction code not found"

def read_source(path):
    wb=load_workbook(path,data_only=True)
    snippets=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1,max_row=min(19,ws.max_row),values_only=True):
            snippets.append(" ".join(str(v) for v in row if v is not None))
    text="\n".join(snippets)
    candidates=[]
    for ws in wb.worksheets:
        hr,headers,score=find_header(ws); mapping=map_columns(headers)
        if score>=2 and len(mapping)>=2:
            names=[str(v).strip() if v is not None else "" for v in headers]
            data=[]
            for ri,vals in enumerate(ws.iter_rows(min_row=hr+1,values_only=True),start=hr+1):
                if not vals or not any(v is not None and str(v).strip() for v in vals): continue
                rec={names[i] if i<len(names) and names[i] else f"Col{i+1}":(vals[i] if i<len(vals) else None) for i in range(len(names))}
                rec["_source_row"]=ri; data.append(rec)
            candidates.append((score,len(data),ws.title,names,mapping,pd.DataFrame(data)))
    if not candidates: raise ValueError("Could not identify a tabular SAP header with columns such as Pos, Material, Bezeichnung, Menge, HkGr or Gesamtwert.")
    candidates.sort(key=lambda x:(x[0],x[1]),reverse=True)
    _,_,sheet,headers,mapping,df=candidates[0]
    tx,mode,evidence=detect_tx(text+" "+" ".join(wb.sheetnames),headers)
    metadata={}
    labels={
      "Material":["material number","materialnummer","matnr"],
      "Werk":["werk","plant"],"Currency":["currency","währung","waehrung"],
      "Costing Variant":["costing variant","kalkulationsvariante"],
      "Base Qty":["base quantity","base qty","grundmenge","basismenge"],
      "Costing Date":["costing date","kalkulationsdatum","kalkulationsstichtag"],
      "Material Description":["material description","materialbeschreibung"]
    }
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1,max_row=min(19,ws.max_row),values_only=True):
            for i,cell in enumerate(row):
                if cell is None: continue
                s=str(cell).strip(); n=norm(s)
                for key,ls in labels.items():
                    for label in ls:
                        ln=norm(label)
                        if n==ln or n.startswith(ln+":") or n.startswith(ln+"="):
                            m=re.match(r"(?i)^"+re.escape(label)+r"\s*[:=]\s*(.+)$",s)
                            val=m.group(1).strip() if m else (str(row[i+1]).strip() if i+1<len(row) and row[i+1] is not None else "")
                            if val and key not in metadata: metadata[key]=val
    return wb,df,headers,mapping,metadata,tx,mode,evidence,sheet,text

def canonicalize(df, mode):
    out=pd.DataFrame(index=df.index)
    for key in set(ASSEMBLY+BOM+OPS+SUPPLIERS+["Ressource","ArbPlatz","STD_Hours","Kostenstelle","Leistungsart"]):
        match=None
        for col in df.columns:
            nc=norm(col)
            aliases=ALIASES.get(key,[key])
            if nc in [norm(a) for a in aliases] or any(len(norm(a))>3 and norm(a) in nc for a in aliases):
                match=col; break
        out[key]=df[match] if match is not None else None
    for col in out.columns:
        if col not in {"Menge","Einzelwert_EUR","Gesamtwert_EUR","STD_Hours","BOM_Level"}:
            out[col]=out[col].map(lambda x:x.strip() if isinstance(x,str) else x)
    for col in ["Menge","Einzelwert_EUR","Gesamtwert_EUR","STD_Hours","BOM_Level"]:
        out[col]=out[col].map(parse_number)
    for idx,val in out["Ressource"].items():
        if isinstance(val,str):
            bits=re.split(r"\s*[|;/]\s*",val.strip())
            if len(bits)>=2:
                if not out.at[idx,"Werk"]: out.at[idx,"Werk"]=bits[0]
                if not out.at[idx,"Material_No"]: out.at[idx,"Material_No"]=bits[1]
                if len(bits)>=3 and not out.at[idx,"Operation"]: out.at[idx,"Operation"]=bits[2]
    for idx,row in out.iterrows():
        txt=" ".join(str(row.get(c) or "") for c in ["HkGr","Bezeichnung","Ressource","ArbPlatz","Operation"]).lower()
        if any(w in txt for w in ["overhead","gemeinkost","zuschlag"]): rt="Overhead"
        elif any(w in txt for w in ["operation","vorgang","bearbeitung","rüst","ruest","maschine","setup","execution"]) and (row.get("ArbPlatz") or row.get("STD_Hours") is not None or row.get("Operation") or any(w in str(row.get("Bezeichnung") or "").lower() for w in ["bearbeitung","operation","vorgang","setup","rüst","ruest","maschine"])): rt="Operation"
        else: rt="Material"
        out.at[idx,"Row_Type"]=rt
        if rt=="Operation":
            ov=str(row.get("Op_Variant") or "").lower(); desc=str(row.get("Bezeichnung") or "").lower()
            if any(w in ov+desc for w in ["setup","rüst","ruest","einrichten"]): opv="Setup"
            elif any(w in ov+desc for w in ["execution","exec","bearbeitung","laufzeit","fertigen"]): opv="Execution"
            else: opv=""
            out.at[idx,"Op_Variant"]=opv
    if mode=="flat": out["BOM_Level"]=1
    for idx,row in out.iterrows():
        pos=str(row.get("Pos") or "")
        if pd.isna(row.get("BOM_Level")):
            out.at[idx,"BOM_Level"]=len(re.findall(r"[./-]",pos))+2 if re.match(r"^\d+(?:[./-]\d+)+$",pos) else None
        level=out.at[idx,"BOM_Level"]
        out.at[idx,"Level_Label"]=f"L{int(level)} – "+("Sub-Assembly" if int(level)==2 else "Component") if pd.notna(level) else ""
        out.at[idx,"Parent_Path"]=" > ".join(re.split(r"[./-]",pos)[:-1]) if re.match(r"^\d+(?:[./-]\d+)+$",pos) else ""
    mask=out[["Pos","Material","Bezeichnung","Ressource","Gesamtwert_EUR"]].apply(lambda row:any(x is not None and not (isinstance(x,float) and math.isnan(x)) and str(x).strip() for x in row),axis=1)
    return out[mask].copy()

def frame_for(df, cols):
    x=df.copy()
    for c in cols:
        if c not in x: x[c]=None
    return x[cols]

def generate(src,dest):
    _,raw,headers,_,metadata,tx,mode,evidence,source_sheet,text=read_source(src)
    if tx=="CK13N": mode="full_bom"
    elif tx=="CK53N": mode="flat"
    elif tx=="ZCO6": mode="period_comparison"
    out=canonicalize(raw,mode)
    assembly=frame_for(out,ASSEMBLY)
    bom=frame_for(out,BOM)
    ops=frame_for(out[(out.Row_Type=="Operation")&(out.Op_Variant=="Execution")],OPS+["Kostenstelle","Leistungsart"])
    vendor_present=any(norm(h) in [norm(a) for a in ALIASES["Kreditor"]] or any(norm(a) in norm(h) for a in ALIASES["Kreditor"] if len(norm(a))>3) for h in headers)
    vend=out[out.Kreditor.notna()] if vendor_present else out.iloc[0:0]
    suppliers=frame_for(vend,SUPPLIERS)
    pb=assembly.copy(); pb["_grp"]=pb.HkGr.map(lambda x:str(x).strip() if x is not None and str(x).strip() else "(unknown)")
    pb["_total"]=pd.to_numeric(pb.Gesamtwert_EUR,errors="coerce")
    pivot=pb.groupby("_grp",dropna=False).agg(Count=("Material","size"),Total_EUR=("_total",lambda x:x.sum(min_count=1))).reset_index().rename(columns={"_grp":"HkGr"})
    total=float(pd.to_numeric(assembly.Gesamtwert_EUR,errors="coerce").sum()) if pd.to_numeric(assembly.Gesamtwert_EUR,errors="coerce").notna().any() else None
    denom=total or 0
    pivot["Share_%"]=pivot.Total_EUR.apply(lambda x:float(x)/denom if denom and pd.notna(x) else None)
    pivot=pivot.sort_values("Total_EUR",ascending=False,na_position="last")[PIVOT]
    currency=metadata.get("Currency","EUR")
    if len(str(currency))>10 or norm(currency).startswith(("pos","hkg","material","bezeichnung","menge")): currency="EUR"
    wb=Workbook(); cover=wb.active; cover.title="Cover"
    for s in SHEETS[1:]: wb.create_sheet(s)
    coverdata=[
      ["LHT_HK_Formatter — SAP Controlling Cost Report"],["Source file",os.path.basename(src)],["Source worksheet",source_sheet],
      ["Source transaction",tx],["Detection evidence",evidence],["Mode",mode],["Material",metadata.get("Material","")],
      ["Werk",metadata.get("Werk","")],["Costing Variant",metadata.get("Costing Variant","")],["Costing Date",metadata.get("Costing Date","")],
      ["Currency",currency],["Base Qty",metadata.get("Base Qty","")],["Total HK",total if total is not None else "N/A"],
      ["Cost summary basis","Sum of detectable line-level Gesamtwert_EUR. Review whether parent and component lines overlap before relying on total."],
      ["Kaufteile %","N/A"],["Bearbeitung %","N/A"],["Overhead %","N/A"],
      ["Vendor note","" if vendor_present else "Vendor data not available in this export"],
      ["Sheet guide","Cover: metadata/validation. Assembly: item summary. BOM_Detail: hierarchy/resource detail. Pivot_Assembly: HkGr summary. Operations: execution operations. Suppliers: vendor reference."],
      ["Rows — Assembly",len(assembly)],["Rows — BOM_Detail",len(bom)],["Rows — Operations",len(ops)],["Rows — Suppliers",len(suppliers)],["Rows — Pivot_Assembly",len(pivot)]
    ]
    if tx=="CK53N": coverdata.append(["Source note","Source: CK53N — flat itemization mode"])
    if tx=="CK13N": coverdata.append(["Source note","Source: CK13N — full BOM mode"])
    if mode=="period_comparison": coverdata.append(["Period comparison","Period comparison mode detected. Confirm comparable period columns before calculating variances; this layout-specific variance analysis is not auto-generated in this baseline implementation."])
    if not vendor_present: coverdata.append(["Vendor warning","Vendor data not available in this export"])
    if mode=="text paste": coverdata.append(["Source note","Source: text paste input"])
    for row in coverdata: cover.append(row)
    specs=[("Assembly",assembly,"Tbl_Assembly"),("BOM_Detail",bom,"Tbl_BOM"),("Pivot_Assembly",pivot,None),("Operations",ops,"Tbl_Operations"),("Suppliers",suppliers,"Tbl_Suppliers")]
    for sname,frame,table_name in specs:
        ws=wb[sname]; ws.append(list(frame.columns))
        for _,row in frame.iterrows():
            values=[]
            for c,v in row.items():
                if pd.isna(v): v=None
                if c in {"Menge","Einzelwert_EUR","Gesamtwert_EUR","STD_Hours","BOM_Level","Count","Total_EUR"} and v is not None: v=parse_number(v)
                values.append(v)
            ws.append(values)
        if table_name:
            tab=Table(displayName=table_name,ref=f"A1:{get_column_letter(len(frame.columns))}{max(2,ws.max_row)}")
            tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False)
            ws.add_table(tab)
        ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor=NAVY); cell.font=Font(bold=True,color=WHITE)
            cell.alignment=Alignment(wrap_text=True,vertical="center")
        for ci in range(1,ws.max_column+1):
            letter=get_column_letter(ci)
            values=[len(str(ws.cell(r,ci).value or "")) for r in range(1,min(ws.max_row,200)+1)]
            ws.column_dimensions[letter].width=min(40,max(12,max(values or [12])+2))
            header=norm(ws.cell(1,ci).value)
            for ri in range(2,ws.max_row+1):
                cell=ws.cell(ri,ci); cell.alignment=Alignment(vertical="top",wrap_text=header in {"bezeichnung","parent_path","ressource"})
                if "gesamtwert" in header or "einzelwert" in header or header=="total_eur": cell.number_format="#,##0.00"
                if header in {"menge","std_hours"}: cell.number_format="#,##0.000"
                if header=="share_%": cell.number_format="0.00%"
        for ci,col in enumerate(frame.columns,1):
            if len(frame)>0 and frame[col].isna().all():
                for ri in range(1,ws.max_row+1): ws.cell(ri,ci).fill=PatternFill("solid",fgColor=YELLOW)
        if sname=="Assembly" and "Gesamtwert_EUR" in frame:
            ci=list(frame.columns).index("Gesamtwert_EUR")+1; letter=get_column_letter(ci)
            if ws.max_row>=2:
                ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}",CellIsRule(operator="between",formula=["10000.000001","50000"],fill=PatternFill("solid",fgColor=GOLD)))
                ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}",CellIsRule(operator="greaterThan",formula=["50000"],fill=PatternFill("solid",fgColor=AMBER)))
        if sname=="BOM_Detail":
            typecol=list(frame.columns).index("Row_Type")+1
            for ri in range(2,ws.max_row+1):
                typ=str(ws.cell(ri,typecol).value or "").lower()
                color=AMBER if typ=="operation" else LIGHT_GREY if typ=="overhead" else (LIGHT_BLUE if ri%2==0 else WHITE)
                for ci in range(1,ws.max_column+1): ws.cell(ri,ci).fill=PatternFill("solid",fgColor=color)
    cover.merge_cells("A1:B1"); cover["A1"].fill=PatternFill("solid",fgColor=NAVY); cover["A1"].font=Font(bold=True,color=WHITE,size=15)
    cover.column_dimensions["A"].width=32; cover.column_dimensions["B"].width=88; cover.freeze_panes="A2"; cover.sheet_view.showGridLines=False
    for ri in range(2,cover.max_row+1):
        cover.cell(ri,1).font=Font(bold=True,color=NAVY); cover.cell(ri,1).fill=PatternFill("solid",fgColor=LIGHT_BLUE)
        cover.cell(ri,2).alignment=Alignment(wrap_text=True,vertical="top")
    opsws=wb["Operations"]; hcol=OPS.index("STD_Hours")+1; r=opsws.max_row+2
    opsws.cell(r,1,"Total STD_Hours").font=Font(bold=True,color=WHITE); opsws.cell(r,1).fill=PatternFill("solid",fgColor=NAVY)
    opsws.cell(r,hcol,f"=SUM({get_column_letter(hcol)}2:{get_column_letter(hcol)}{max(2,opsws.max_row)})").number_format="#,##0.000"
    wb.save(dest)
    # Re-open to verify sheet names, tables, and required headers.
    check=load_workbook(dest,data_only=False)
    errors=[]
    if check.sheetnames!=SHEETS: errors.append("Sheet names/order")
    for s,t in [("Assembly","Tbl_Assembly"),("BOM_Detail","Tbl_BOM"),("Operations","Tbl_Operations"),("Suppliers","Tbl_Suppliers")]:
        if t not in check[s].tables: errors.append(f"Missing {t}")
    for s,cols in [("Assembly",ASSEMBLY),("BOM_Detail",BOM),("Pivot_Assembly",PIVOT),("Operations",OPS+["Kostenstelle","Leistungsart"]),("Suppliers",SUPPLIERS)]:
        if [check[s].cell(1,c).value for c in range(1,len(cols)+1)]!=cols: errors.append(f"Column contract: {s}")
    for ri in range(1,check["Cover"].max_row+1):
        if check["Cover"].cell(ri,1).value=="Validation status": check["Cover"].cell(ri,2,"FAIL" if errors else "PASS")
    check["Cover"].append(["Validation status","FAIL" if errors else "PASS"])
    check["Cover"].append(["Validation detail","; ".join(errors) if errors else "Sheet, table and required header structure checks passed. Review source-specific mapping and period comparisons before operational use."])
    check.save(dest)
    return {"output":dest,"transaction":tx,"mode":mode,"evidence":evidence,"rows_assembly":len(assembly),"rows_bom":len(bom),"rows_operations":len(ops),"rows_suppliers":len(suppliers),"rows_pivot":len(pivot),"total_hk":total,"currency":currency,"validation":"PASS" if not errors else "FAIL","errors":errors}

def main():
    p=argparse.ArgumentParser()
    p.add_argument("input"); p.add_argument("-o","--output")
    a=p.parse_args()
    if not os.path.isfile(a.input): raise SystemExit("Input file not found: "+a.input)
    out=a.output or str(Path(a.input).with_name(f"LHT_HK_Unknown_Unknown_{date.today().isoformat()}_Formatted.xlsx"))
    print(json.dumps(generate(a.input,out),ensure_ascii=False,indent=2,default=str))
if __name__=="__main__": main()
